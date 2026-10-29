from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# ── Page margins ──
section = doc.sections[0]
section.page_width  = Inches(8.27)
section.page_height = Inches(11.69)
section.left_margin   = Inches(1.1)
section.right_margin  = Inches(1.1)
section.top_margin    = Inches(1.0)
section.bottom_margin = Inches(1.0)

# ── Styles helper ──
def style_para(para, size, bold=False, color=None, space_before=0, space_after=6, align=None):
    para.paragraph_format.space_before = Pt(space_before)
    para.paragraph_format.space_after  = Pt(space_after)
    for run in para.runs:
        run.font.size = Pt(size)
        run.font.bold = bold
        if color:
            run.font.color.rgb = RGBColor(*color)
    if align:
        para.alignment = align

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    if level == 1:
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after  = Pt(6)
        # Bottom border
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '4')
        bottom.set(qn('w:color'), '000000')
        pBdr.append(bottom)
        pPr.append(pBdr)
    elif level == 2:
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after  = Pt(4)
    return p

def add_body(doc, text, size=10.5, indent=False):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(5)
    if indent:
        p.paragraph_format.left_indent = Cm(0.5)
    for run in p.runs:
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(30, 30, 30)
    return p

def add_bullet(doc, text, size=10.5):
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(30, 30, 30)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after  = Pt(3)
    p.paragraph_format.left_indent  = Cm(0.5)
    return p

def add_table(doc, headers, rows):
    table = doc.add_table(rows=1+len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Header row
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = h
        run = cell.paragraphs[0].runs[0]
        run.font.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)
        # Black background
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'),   'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'),  '1A1A1A')
        tcPr.append(shd)
        cell.paragraphs[0].paragraph_format.space_before = Pt(3)
        cell.paragraphs[0].paragraph_format.space_after  = Pt(3)
    # Data rows
    for ri, row_data in enumerate(rows):
        row = table.rows[ri + 1]
        fill = 'F7F7F7' if ri % 2 == 0 else 'FFFFFF'
        for ci, val in enumerate(row_data):
            cell = row.cells[ci]
            cell.text = val
            run = cell.paragraphs[0].runs[0]
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(30, 30, 30)
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'),   'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'),  fill)
            tcPr.append(shd)
            cell.paragraphs[0].paragraph_format.space_before = Pt(2)
            cell.paragraphs[0].paragraph_format.space_after  = Pt(2)
    return table

# ══════════════════════════════════════════════
#  COVER / TITLE BLOCK
# ══════════════════════════════════════════════
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = title_p.add_run("DocPilot")
r.font.size = Pt(28)
r.font.bold = True
r.font.color.rgb = RGBColor(0, 0, 0)
title_p.paragraph_format.space_before = Pt(0)
title_p.paragraph_format.space_after  = Pt(2)

sub_p = doc.add_paragraph()
rs = sub_p.add_run("Testing & Validation Report")
rs.font.size = Pt(14)
rs.font.bold = False
rs.font.color.rgb = RGBColor(80, 80, 80)
sub_p.paragraph_format.space_after = Pt(2)

meta_p = doc.add_paragraph()
rm = meta_p.add_run("Intelligent Document Intake & Automated Routing System  |  Version 1.0  |  October 2026")
rm.font.size = Pt(9)
rm.font.color.rgb = RGBColor(120, 120, 120)
meta_p.paragraph_format.space_after = Pt(14)

# Divider line
div = doc.add_paragraph()
div.paragraph_format.space_before = Pt(0)
div.paragraph_format.space_after  = Pt(14)
pPr = div._p.get_or_add_pPr()
pBdr = OxmlElement('w:pBdr')
bottom = OxmlElement('w:bottom')
bottom.set(qn('w:val'),   'single')
bottom.set(qn('w:sz'),    '12')
bottom.set(qn('w:space'), '1')
bottom.set(qn('w:color'), '000000')
pBdr.append(bottom)
pPr.append(pBdr)

# ══════════════════════════════════════════════
#  1. PURPOSE
# ══════════════════════════════════════════════
add_heading(doc, "1.  Purpose & Scope")
add_body(doc,
    "This document covers the complete testing and validation process carried out for DocPilot — "
    "an AI-powered document filing system. The goal was to make sure every part of the system "
    "works correctly, handles errors gracefully, and meets the requirements defined at the start "
    "of the project. Testing was done in three layers: unit, integration, and full system testing, "
    "followed by performance and basic security checks.")

# ══════════════════════════════════════════════
#  2. UNIT TESTING
# ══════════════════════════════════════════════
add_heading(doc, "2.  Unit Testing")
add_body(doc,
    "Unit testing checks each module individually to confirm it produces the correct output "
    "for a given input, without depending on other parts of the system.")

add_heading(doc, "Modules Tested", level=2)

add_table(doc,
    ["Module", "What Was Tested", "Result"],
    [
        ["extractor.py",      "Correct text extraction from PDF files; handles blank/corrupt PDFs without crashing", "✓ Pass"],
        ["classifier.py",     "Keyword matching returns correct category (Invoice, Contract, Lease, Receipt, Other)", "✓ Pass"],
        ["ai_classifier.py",  "Gemini API returns type + confidence + reasoning; OpenAI fallback triggers when Gemini key is missing; keyword fallback runs offline", "✓ Pass"],
        ["filer.py",          "Correct folder name generated for any doc type; new folders created automatically; duplicate file names handled", "✓ Pass"],
        ["database.py",       "Insert, read, and update operations; schema migration runs without errors on older DB", "✓ Pass"],
        ["config.py",         "API keys load from .env file; missing keys default to empty string without crash", "✓ Pass"],
    ]
)
doc.add_paragraph().paragraph_format.space_after = Pt(6)

add_body(doc,
    "All unit tests were run using Python's built-in test scripts (test_pipeline.py, test_app.py). "
    "Every module passed with no exceptions under normal inputs. Edge cases like empty PDFs, "
    "missing API keys, and unknown document types were also tested.")

# ══════════════════════════════════════════════
#  3. INTEGRATION TESTING
# ══════════════════════════════════════════════
add_heading(doc, "3.  Integration Testing")
add_body(doc,
    "Integration testing checks that the modules work correctly when connected together "
    "as a pipeline. The full flow — Upload → Extract → Classify → File → Save to DB — was "
    "tested end-to-end.")

add_heading(doc, "Test Scenarios", level=2)

add_bullet(doc, "Uploaded a valid Invoice PDF → correctly classified as Invoice → filed in Invoices/ folder → DB record created with correct type, confidence, and provider.")
add_bullet(doc, "Uploaded a Contract PDF → confidence above 85% → auto-filed without human review → appeared in Filed Log.")
add_bullet(doc, "Uploaded an ambiguous document → confidence below 85% → sent to Review Queue → user confirmed category → manually filed to correct folder.")
add_bullet(doc, "Uploaded a duplicate file (same content) → SHA-256 hash matched → system blocked duplicate and showed error message.")
add_bullet(doc, "Saved a Gemini API key via Settings page → key written to .env → AI engine activated without restarting server.")
add_bullet(doc, "Removed API key → system automatically fell back to keyword classifier → offline classification still worked.")
add_bullet(doc, "Clicked a folder in sidebar → folder_view page loaded with correct file list and DB records.")

doc.add_paragraph().paragraph_format.space_after = Pt(4)
add_body(doc, "All integration scenarios passed. No data loss or misrouting was found during these tests.")

# ══════════════════════════════════════════════
#  4. SYSTEM TESTING
# ══════════════════════════════════════════════
add_heading(doc, "4.  System Testing")
add_body(doc,
    "System testing validates the entire application as a single unit — from the browser UI "
    "to the database — to make sure it meets the original requirements.")

add_heading(doc, "Requirements Validation", level=2)

add_table(doc,
    ["Requirement", "Test Method", "Status"],
    [
        ["Upload PDF and classify automatically",         "Uploaded 5 different PDF types",                          "✓ Met"],
        ["Auto-file docs with confidence ≥ 85%",          "Verified 10 uploads — all high-confidence docs filed",    "✓ Met"],
        ["Hold low-confidence docs for human review",     "Docs under 85% appeared in Review Queue",                "✓ Met"],
        ["Create new folders for unknown document types", "Uploaded a 'Bank Statement' PDF — new folder created",    "✓ Met"],
        ["AI classification with Gemini / OpenAI",        "Both keys tested individually; results verified",         "✓ Met"],
        ["Keyword fallback when no API key",              "Removed keys — keyword classifier ran offline",           "✓ Met"],
        ["Duplicate file detection",                      "Same file uploaded twice — second blocked by hash check", "✓ Met"],
        ["Per-folder view of filed documents",            "Clicked each sidebar folder — file list displayed",       "✓ Met"],
        ["Settings page saves API key to .env",           "Key saved, server picked up without full restart",        "✓ Met"],
        ["Folder watcher detects files dropped in folder","Dropped PDF in watched folder — auto-classified",         "✓ Met"],
    ]
)
doc.add_paragraph().paragraph_format.space_after = Pt(6)

# ══════════════════════════════════════════════
#  5. DEFECTS FOUND & FIXED
# ══════════════════════════════════════════════
add_heading(doc, "5.  Defects Found & Fixed")

add_table(doc,
    ["#", "Defect", "Impact", "Fix Applied"],
    [
        ["D-01", "Flask debug reloader conflicted with Watchdog folder monitor — server port never bound correctly", "Critical — app unreachable", "Added use_reloader=False to app.run()"],
        ["D-02", "folder_view.html template missing — clicking any sidebar folder caused HTTP 500 error",           "High — feature broken",     "Created folder_view.html with full DB + disk view"],
        ["D-03", "Old DB schema missing reasoning, entities, provider columns — schema migration failed silently",   "Medium — data not saved",   "ALTER TABLE with IF NOT EXISTS check added in init_db()"],
        ["D-04", "Masked API key (•••) was accidentally overwriting .env with bullet characters",                   "High — broke API key",      "is_masked() check added before writing .env"],
        ["D-05", "Duplicate file uploaded without error — same document filed twice",                               "Medium — duplicate records","SHA-256 hash stored and checked before every upload"],
    ]
)
doc.add_paragraph().paragraph_format.space_after = Pt(6)

# ══════════════════════════════════════════════
#  6. PERFORMANCE
# ══════════════════════════════════════════════
add_heading(doc, "6.  Performance Testing")
add_body(doc,
    "Basic performance checks were done to ensure the system responds within acceptable time under "
    "normal use (single-user desktop application).")

add_table(doc,
    ["Test", "Result", "Acceptable?"],
    [
        ["PDF text extraction (5-page PDF)",       "< 1 second",    "✓ Yes"],
        ["Gemini AI classification (API call)",    "2 – 4 seconds", "✓ Yes"],
        ["Keyword classification (offline)",       "< 0.1 second",  "✓ Yes"],
        ["File move to destination folder",        "< 0.2 second",  "✓ Yes"],
        ["Dashboard load (8 recent docs)",         "< 0.5 second",  "✓ Yes"],
        ["Folder view load (20 files)",            "< 0.5 second",  "✓ Yes"],
    ]
)
doc.add_paragraph().paragraph_format.space_after = Pt(6)

add_body(doc,
    "The only slowness observed was the Gemini/OpenAI API call (2-4 seconds), which depends on "
    "internet speed. This is expected and acceptable for a document filing system.")

# ══════════════════════════════════════════════
#  7. SECURITY
# ══════════════════════════════════════════════
add_heading(doc, "7.  Security Checks")

add_bullet(doc, "API keys stored in .env file — never hardcoded in source code or shown in full on the Settings page (masked with •••).")
add_bullet(doc, "File upload restricted to PDF format only — non-PDF uploads are rejected with an error message.")
add_bullet(doc, "Folder path in /open-folder route sanitised — only alphanumeric, underscore, and hyphen characters allowed to prevent directory traversal.")
add_bullet(doc, "SQLite database stored locally — no external network access; no user authentication needed for single-user desktop use.")
add_bullet(doc, ".gitignore excludes .env, database, and uploaded files — sensitive data not accidentally committed to version control.")

# ══════════════════════════════════════════════
#  8. VALIDATION SUMMARY
# ══════════════════════════════════════════════
add_heading(doc, "8.  Validation Summary")
add_body(doc,
    "DocPilot has been fully validated against all stated requirements. All 10 core features "
    "passed system testing. Five defects identified during development were fixed before final "
    "release. Performance is well within acceptable limits for a desktop filing tool. "
    "Basic security best practices are in place.")

add_body(doc,
    "The system is stable, reliable, and ready for use. Future versions will add email intake, "
    "multi-file batch upload, and deadline extraction for contract documents.")

# ══════════════════════════════════════════════
#  Footer line
# ══════════════════════════════════════════════
doc.add_paragraph()
foot = doc.add_paragraph()
foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
rf = foot.add_run("DocPilot v1.0  ·  Testing & Validation Report  ·  Confidential")
rf.font.size  = Pt(8)
rf.font.color.rgb = RGBColor(160, 160, 160)

# ── Save ──
out = r"e:\Macro project\Testing_and_Validation.docx"
doc.save(out)
print(f"Saved: {out}")
